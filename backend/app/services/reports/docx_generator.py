import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from app.schemas.domain import InspectionReportModel

class DocxReportGenerator:
    @staticmethod
    def generate_docx(report_data: InspectionReportModel) -> bytes:
        doc = Document()

        # Set page margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        # Title / Branding
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("LM-TRACE")
        run.font.name = "Arial"
        run.font.size = Pt(22)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 37, 55)  # Navy

        subtitle = doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = subtitle.add_run("AI-Assisted Legal Metrology Inspection & Compliance Report")
        sub_run.font.name = "Arial"
        sub_run.font.size = Pt(13)
        sub_run.font.color.rgb = RGBColor(30, 64, 175)  # Royal Blue

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

        # 1. Inspection Metadata Table
        meta_table = doc.add_table(rows=5, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_data = [
            ("Inspection ID / Code:", f"{report_data.inspection_code} (ID: {report_data.inspection_id})"),
            ("Inspection Date & Time:", report_data.inspection_date),
            ("Inspector Details:", f"{report_data.inspector_name} (Officer ID: {report_data.officer_id})"),
            ("Inspection Location:", report_data.location),
            ("Establishment / Seller:", f"{report_data.seller_name or 'N/A'} ({report_data.business_name or 'N/A'})")
        ]
        for idx, (lbl, val) in enumerate(meta_data):
            row = meta_table.rows[idx]
            r0 = row.cells[0].paragraphs[0].add_run(lbl)
            r0.font.bold = True
            r0.font.size = Pt(9.5)
            r1 = row.cells[1].paragraphs[0].add_run(str(val))
            r1.font.size = Pt(9.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 2. Product Summary & Overall Assessment
        doc.add_heading("1. Product Information & Overall Assessment", level=2)
        p_prod = doc.add_paragraph()
        p_prod.add_run(f"Product Name: ").bold = True
        p_prod.add_run(f"{report_data.product_name} | ")
        p_prod.add_run(f"Brand: ").bold = True
        p_prod.add_run(f"{report_data.brand or 'N/A'} | ")
        p_prod.add_run(f"Category: ").bold = True
        p_prod.add_run(f"{report_data.category}\n")
        p_prod.add_run(f"Overall Compliance Status: ").bold = True
        status_run = p_prod.add_run(f"{report_data.overall_status} ")
        status_run.bold = True
        p_prod.add_run(f"| Analytical Assessment Score: ").bold = True
        p_prod.add_run(f"{report_data.score} / 100")

        # 3. Declarations Matrix Table (Dual AI vs Inspector Verified Values)
        doc.add_heading("2. Statutory Declarations & Verification Matrix", level=2)
        dec_headers = ["Declaration", "AI Extracted Value", "Inspector Verified Value", "Correctness", "Status"]
        dec_table = doc.add_table(rows=1, cols=5)
        dec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr_cells = dec_table.rows[0].cells
        for i, header_text in enumerate(dec_headers):
            h_run = hdr_cells[i].paragraphs[0].add_run(header_text)
            h_run.font.bold = True
            h_run.font.size = Pt(9)

        for dec in report_data.declarations:
            row_cells = dec_table.add_row().cells
            row_cells[0].paragraphs[0].add_run(dec.get("declaration", "")).font.size = Pt(8.5)
            row_cells[1].paragraphs[0].add_run(dec.get("ai_value") or "Missing").font.size = Pt(8.5)
            row_cells[2].paragraphs[0].add_run(dec.get("verified_value") or dec.get("ai_value") or "Unverified").font.size = Pt(8.5)
            row_cells[3].paragraphs[0].add_run(dec.get("correctness", "VALID")).font.size = Pt(8.5)
            s_run = row_cells[4].paragraphs[0].add_run(dec.get("final_check", "PASS"))
            s_run.font.size = Pt(8.5)
            s_run.font.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 4. Compliance Checks Matrix
        doc.add_heading("3. Legal Metrology Compliance Checks", level=2)
        chk_headers = ["Rule Code", "Check Description", "Input / Measurement", "Rule Standard", "Result"]
        chk_table = doc.add_table(rows=1, cols=5)
        chk_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(chk_headers):
            chk_table.rows[0].cells[i].paragraphs[0].add_run(h).font.bold = True

        for chk in report_data.compliance_checks:
            r = chk_table.add_row().cells
            r[0].paragraphs[0].add_run(chk.get("rule_code", "")).font.size = Pt(8.5)
            r[1].paragraphs[0].add_run(chk.get("check_type", "")).font.size = Pt(8.5)
            r[2].paragraphs[0].add_run(str(chk.get("input_value", ""))).font.size = Pt(8.5)
            r[3].paragraphs[0].add_run(chk.get("expected_condition", "")).font.size = Pt(8.5)
            res_run = r[4].paragraphs[0].add_run(chk.get("result", ""))
            res_run.font.size = Pt(8.5)
            res_run.font.bold = True

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        # 5. Inspector Remarks & Disclaimers
        doc.add_heading("4. Inspector Remarks & Regulatory Disclaimer", level=2)
        p_rem = doc.add_paragraph()
        p_rem.add_run("Inspector Observations & Orders:\n").bold = True
        p_rem.add_run(report_data.inspector_remarks or "Inspection completed with visual evidence and digital audit trail recorded.")

        p_disc = doc.add_paragraph()
        p_disc.paragraph_format.space_before = Pt(12)
        d_run = p_disc.add_run(f"LEGAL DISCLAIMER: {report_data.disclaimer}")
        d_run.font.size = Pt(8)
        d_run.font.italic = True
        d_run.font.color.rgb = RGBColor(100, 116, 139)

        # Signature Block
        sig_p = doc.add_paragraph()
        sig_p.paragraph_format.space_before = Pt(20)
        sig_p.add_run("____________________________\n").bold = True
        sig_p.add_run(f"Verified & Authorized by: {report_data.inspector_name}\n")
        sig_p.add_run(f"Officer ID: {report_data.officer_id} | Date: {report_data.inspection_date}\n")
        sig_p.add_run("Legal Metrology Department, Government of India")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

docx_report_generator = DocxReportGenerator()
